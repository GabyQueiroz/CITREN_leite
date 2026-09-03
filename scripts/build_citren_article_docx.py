from pathlib import Path
import re

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\gabri\Documents\UEPG\MestradoAndressa")
TEMPLATE = Path(r"C:\Users\gabri\Downloads\formato ejemplo CITREN (1).docx")
OUT = ROOT / "doc_work" / "articulo_CITREN_leite_materiales_resultados.docx"

DATA = ROOT / "article_outputs" / "data"
TABLES = ROOT / "article_outputs" / "tables"
FIGS = ROOT / "article_outputs" / "figures"

METRICS = TABLES / "metricas_modelos.csv"
IMPORTANCE = TABLES / "importancia_variables.csv"
PREDICTIONS = TABLES / "predicciones_prueba.csv"
BASE_INTEGRADA = DATA / "base_integrada_leite_agroclima.csv"
CLIMA_ANUAL = DATA / "clima_castro_a819_anual.csv"
SCRIPT = ROOT / "scripts" / "build_article_models.py"


def norm(text):
    return re.sub(r"\s+", " ", text.strip().upper())


def clear_paragraph(paragraph):
    for run in list(paragraph.runs):
        paragraph._element.remove(run._element)


def set_font(run, size=10, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def format_paragraph(paragraph, size=10, bold=False, align=None, space_after=6, first_line=True):
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.0
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(0.75)
    else:
        paragraph.paragraph_format.first_line_indent = Cm(0)
    for run in paragraph.runs:
        set_font(run, size=size, bold=bold)


def insert_p_after(paragraph, text="", size=10, bold=False, italic=False, align=None, space_after=6, first_line=True):
    new = OxmlElement("w:p")
    paragraph._p.addnext(new)
    p = paragraph._parent.add_paragraph()
    p._p = new
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, italic=italic)
    format_paragraph(p, size=size, bold=bold, align=align, space_after=space_after, first_line=first_line)
    return p


def insert_table_after(paragraph, data, widths=None, font_size=8):
    table = paragraph._parent.add_table(rows=len(data), cols=len(data[0]), width=Cm(16))
    paragraph._p.addnext(table._tbl)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    table.autofit = True
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = str(value)
            if widths:
                cell.width = widths[c]
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.first_line_indent = Cm(0)
                for run in p.runs:
                    set_font(run, size=font_size, bold=(r == 0))
            if r == 0:
                tc_pr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "D9EAD3")
                tc_pr.append(shd)
            set_cell_borders(cell)
    p_after = paragraph._parent.add_paragraph()
    table._tbl.addnext(p_after._p)
    format_paragraph(p_after, first_line=False)
    return table, p_after


def set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "7F7F7F")


def insert_picture_after(paragraph, image_path, caption, width_cm=13.2):
    p_img = insert_p_after(paragraph, "", first_line=False, space_after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p_img.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    p_cap = insert_p_after(p_img, caption, size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=8)
    return p_cap


def delete_between(doc, start_idx, end_idx):
    for p in doc.paragraphs[start_idx + 1 : end_idx]:
        p._element.getparent().remove(p._element)


def find_section_indices(doc):
    labels = {
        "INTRODUCCIÓN": None,
        "MATERIALES Y MÉTODOS": None,
        "RESULTADOS Y DISCUSIÓN": None,
        "CONCLUSIÓN": None,
        "REFERENCIAS BIBLIOGRÁFICAS": None,
    }
    for i, p in enumerate(doc.paragraphs):
        t = norm(p.text)
        for label in labels:
            if norm(label) == t:
                labels[label] = i
    return labels


def heading(paragraph, text):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_font(run, size=10, bold=True)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)


def subheading(anchor, text):
    return insert_p_after(anchor, text, size=10, bold=True, first_line=False, space_after=3)


def fmt(x, decimals=2):
    if isinstance(x, str):
        return x
    return f"{x:,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def main():
    doc = Document(TEMPLATE)

    # Keep page geometry from the template when possible.
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT

    labels = find_section_indices(doc)
    if labels["MATERIALES Y MÉTODOS"] is None or labels["RESULTADOS Y DISCUSIÓN"] is None:
        raise RuntimeError("Template section headings were not found.")

    delete_between(doc, labels["MATERIALES Y MÉTODOS"], labels["RESULTADOS Y DISCUSIÓN"])
    labels = find_section_indices(doc)
    delete_between(doc, labels["RESULTADOS Y DISCUSIÓN"], labels["CONCLUSIÓN"])
    labels = find_section_indices(doc)

    for label in ["MATERIALES Y MÉTODOS", "RESULTADOS Y DISCUSIÓN", "CONCLUSIÓN", "REFERENCIAS BIBLIOGRÁFICAS"]:
        if labels[label] is not None:
            heading(doc.paragraphs[labels[label]], label)

    anchor = doc.paragraphs[labels["MATERIALES Y MÉTODOS"]]
    anchor = subheading(anchor, "Diseño metodológico")
    for text in [
        "El estudio fue desarrollado como un experimento computacional de predicción supervisada para estimar la producción anual de leche de los municipios del estado de Paraná, Brasil. La unidad de observación fue el municipio-año y la variable dependiente fue la producción municipal de leche, expresada en mil litros. La propuesta metodológica se basó en la integración de datos agropecuarios públicos con variables agroclimáticas anuales, seguida por la comparación entre un modelo base temporal y dos algoritmos de aprendizaje automático.",
        "El flujo experimental fue organizado en cinco etapas: recopilación de bases públicas; limpieza y estandarización de registros; construcción de variables temporales y climáticas; entrenamiento y ajuste de algoritmos; y evaluación comparativa mediante métricas de regresión. Esta organización permite que el experimento sea reproducible y que cada resultado pueda ser rastreado hasta el archivo utilizado en el procesamiento.",
        "Aunque se analizaron otras bases públicas para etapas futuras de investigación, como MapBiomas, TOPODATA, CAR, SIGEF, red vial, CHIRPS, ANA/HidroWeb e hidrografía, estas no fueron utilizadas en los modelos presentados en este artículo. La decisión metodológica se justifica porque el objetivo del artículo de congreso fue producir una versión más directa, basada en datos tabulares públicos, enfocada en predicción municipal anual y no en simulación espacial de propiedades, rutas o logística de recolección.",
    ]:
        anchor = insert_p_after(anchor, text)

    anchor = subheading(anchor, "Fuentes públicas de datos")
    source_rows = [
        ["Base", "Institución", "Enlace", "Uso en el artículo"],
        ["Producción de leche municipal", "IBGE/SIDRA - Pesquisa da Pecuária Municipal", "https://sidra.ibge.gov.br/tabela/74", "Variable respuesta anual por municipio, en mil litros."],
        ["Datos meteorológicos históricos", "INMET - estación automática A819 Castro/PR", "https://portal.inmet.gov.br/dadoshistoricos", "Precipitación, temperatura, humedad, radiación y viento agregados por año."],
        ["Repositorio del experimento", "GitHub", "https://github.com/GabyQueiroz/CITREN_leite", "Código, base integrada, tablas, figuras y documento final."],
    ]
    _, anchor = insert_table_after(anchor, source_rows, font_size=7)
    anchor = insert_p_after(anchor, "Tabla 1. Fuentes públicas, enlaces y función de cada fuente en el artículo.", size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)

    anchor = subheading(anchor, "Archivos utilizados y reproducibilidad")
    file_rows = [
        ["Elemento", "Archivo o ruta utilizada", "Función en el estudio"],
        ["Script principal", str(SCRIPT), "Descarga, integración, entrenamiento, evaluación y generación de salidas."],
        ["Base integrada", str(BASE_INTEGRADA), "Panel final municipio-año con producción, variables temporales y clima anual."],
        ["Clima anual", str(CLIMA_ANUAL), "Indicadores anuales agregados a partir de la estación INMET A819."],
        ["Métricas", str(METRICS), "Comparación cuantitativa entre persistencia, Random Forest y XGBoost."],
        ["Predicciones", str(PREDICTIONS), "Valores observados y predichos en el período de prueba."],
        ["Importancia", str(IMPORTANCE), "Importancia de variables del mejor modelo de aprendizaje automático."],
    ]
    _, anchor = insert_table_after(anchor, file_rows, font_size=7)
    anchor = insert_p_after(anchor, "Tabla 2. Archivos locales utilizados para reproducir la base integrada, los algoritmos y los resultados.", size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    anchor = insert_p_after(anchor, "Además del almacenamiento local, los códigos y salidas fueron enviados al repositorio GitHub https://github.com/GabyQueiroz/CITREN_leite. En el repositorio, la carpeta scripts contiene el archivo de procesamiento y modelado; la carpeta data contiene la base integrada y los datos tratados; la carpeta results contiene tablas y figuras; y la carpeta docs contiene el documento del artículo. Esta estructura facilita la auditoría del experimento, la repetición de los modelos y la actualización futura de los datos.")

    anchor = subheading(anchor, "Bases de datos realmente utilizadas")
    for text in [
        "La primera base fue la Pesquisa da Pecuária Municipal, obtenida por medio de la plataforma SIDRA del IBGE, específicamente la tabla 74 para el producto leche. De esta fuente se extrajeron el código del municipio, el nombre del municipio, el año y la cantidad anual de leche producida. La serie disponible comprende registros desde 1974 hasta 2024 para los municipios paranaenses; sin embargo, para el modelado se utilizó el período de 2008 a 2019, por ser el intervalo con compatibilidad entre las variables rezagadas y los datos climáticos tratados.",
        "La segunda base fue compuesta por datos meteorológicos horarios de la estación automática A819, localizada en Castro/PR y disponibilizada por el INMET. Fueron empleados registros de precipitación, temperatura del aire, humedad relativa, radiación solar global y velocidad del viento. Los datos fueron tratados para remover códigos de falla, valores nulos y registros incompatibles, además de convertir separadores decimales, fechas y columnas numéricas.",
        "A partir de los datos horarios del INMET fueron generadas variables anuales: precipitación total acumulada, temperatura media anual, temperatura máxima observada, humedad relativa media, radiación solar media, velocidad media del viento y número aproximado de días secos, definidos como días con precipitación diaria inferior a 1 mm. Estas variables fueron integradas al panel municipal por año, funcionando como una señal agroclimática anual común para el estado.",
        "La opción por una estación de Castro/PR fue adoptada por coherencia con el tema agrícola del estudio y por la disponibilidad de una serie horaria pública procesável. No obstante, esta elección también representa una limitación: el clima anual de una única estación no substituye una matriz espacial completa para todos los municipios. Por eso, los resultados deben ser entendidos como una prueba metodológica reproducible para congreso.",
    ]:
        anchor = insert_p_after(anchor, text)

    anchor = subheading(anchor, "Construcción de la base integrada")
    for text in [
        "La base integrada fue creada a partir de la unión entre la producción municipal de leche y los indicadores climáticos anuales. Inicialmente, los registros de producción fueron ordenados por municipio y año. En seguida, fueron generadas variables de memoria temporal: producción del año anterior, producción de dos años anteriores, media móvil de tres años y crecimiento porcentual observado en el año anterior.",
        "Después de la generación de rezagos, los registros sin historial suficiente fueron eliminados del conjunto de modelado. Esta decisión evita que el algoritmo aprenda con valores imputados artificialmente en los primeros años de cada serie municipal. El resultado fue un panel anual capaz de representar tanto la inercia productiva de cada municipio como variaciones climáticas agregadas.",
    ]:
        anchor = insert_p_after(anchor, text)
    var_rows = [
        ["Grupo", "Variables"],
        ["Identificación", "codigo_municipio, municipio, ano"],
        ["Histórico productivo", "leite_lag_1, leite_lag_2, leite_media_3anos, crescimento_lag_1_pct"],
        ["Clima anual", "precipitacao_anual_mm, temperatura_media_c, temperatura_max_media_c, radiacao_media_kj_m2, umidade_media_pct, vento_medio_ms, dias_secos_aprox"],
        ["Respuesta", "leite_mil_litros"],
    ]
    _, anchor = insert_table_after(anchor, var_rows, font_size=8)
    anchor = insert_p_after(anchor, "Tabla 3. Variables utilizadas en la base final de modelado.", size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)

    anchor = subheading(anchor, "Creación de los algoritmos")
    for text in [
        "Los modelos fueron desarrollados en Python con pandas, scikit-learn, XGBoost y matplotlib. Para las variables numéricas se aplicó imputación por mediana y estandarización. Las variables categóricas asociadas al municipio fueron transformadas mediante one-hot encoding, de modo que cada municipio pudiera ser representado sin imponer una relación ordinal artificial entre códigos.",
        "La variable objetivo fue transformada por log(1 + y) en Random Forest y XGBoost, con retorno posterior a la escala original mediante exp(y) - 1. Esta transformación fue utilizada porque la producción de leche presenta gran diferencia de escala entre municipios: algunos producen volúmenes muy altos, mientras otros poseen producción menor. La transformación logarítmica reduce la influencia excesiva de valores extremos durante el entrenamiento.",
        "La separación de datos respetó la secuencia temporal. Los años hasta 2016 se utilizaron para entrenamiento, validación interna y selección de hiperparámetros; los años de 2017 a 2019 fueron reservados para la prueba final. Esta estrategia evita que observaciones futuras influyan en el entrenamiento y permite una evaluación más realista que una partición aleatoria tradicional.",
        "Se compararon tres enfoques. El primero fue la persistencia lag-1, en la cual la predicción de un municipio en el año t es igual a la producción observada en el año t - 1. El segundo fue Random Forest Regressor, un ensamble de árboles de decisión capaz de modelar relaciones no lineales y interacciones entre variables. El tercero fue XGBoost Regressor, basado en boosting de gradiente, que ajusta árboles secuenciales para reducir errores residuales.",
    ]:
        anchor = insert_p_after(anchor, text)
    hp_rows = [
        ["Modelo", "Hiperparámetros evaluados", "Configuración utilizada"],
        ["Persistencia lag-1", "No aplica", "Predicción igual al año anterior."],
        ["Random Forest", "n_estimators, max_depth, min_samples_leaf", "n_estimators = 400; max_depth = 12; min_samples_leaf = 1."],
        ["XGBoost", "n_estimators, max_depth, learning_rate, subsample, colsample_bytree", "n_estimators = 500; max_depth = 3; learning_rate = 0,05; subsample = 0,90; colsample_bytree = 0,90."],
    ]
    _, anchor = insert_table_after(anchor, hp_rows, font_size=8)
    anchor = insert_p_after(anchor, "Tabla 4. Algoritmos comparados y configuración final de hiperparámetros.", size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    anchor = insert_p_after(anchor, "La selección de hiperparámetros fue realizada con GridSearchCV y validación cruzada temporal TimeSeriesSplit con tres divisiones. La métrica de optimización fue el menor RMSE, ya que esta métrica penaliza con mayor intensidad errores grandes y es relevante cuando se desea evitar estimaciones muy distantes en municipios de alta producción.")
    anchor = insert_p_after(anchor, "Para la evaluación final se calcularon cuatro métricas: MAE, que representa el error absoluto promedio; RMSE, que destaca errores de mayor magnitud; R², que mide la proporción de variabilidad explicada por el modelo; y MAPE, que expresa el error porcentual medio. El uso conjunto de estas métricas permite interpretar tanto la precisión en unidades productivas como la calidad relativa de las predicciones.")

    labels = find_section_indices(doc)
    anchor = doc.paragraphs[labels["RESULTADOS Y DISCUSIÓN"]]
    metrics = pd.read_csv(METRICS)
    met_rows = [["Modelo", "MAE", "RMSE", "R²", "MAPE (%)"]]
    for _, row in metrics.iterrows():
        met_rows.append([row["modelo"], fmt(row["MAE"]), fmt(row["RMSE"]), fmt(row["R2"], 4), fmt(row["MAPE_%"])])
    _, anchor = insert_table_after(anchor, met_rows, font_size=8)
    anchor = insert_p_after(anchor, "Tabla 5. Desempeño de los modelos en la base de prueba de 2017 a 2019.", size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)

    anchor = insert_p_after(anchor, "Los resultados muestran que la producción lechera municipal presenta fuerte dependencia temporal. El modelo de persistencia obtuvo el menor RMSE, con 2.974,01 mil litros, MAE de 1.139,22 mil litros, R² de 0,9800 y MAPE de 20,01%. Este comportamiento indica que, para una serie anual municipal, la producción del año anterior ya contiene gran parte de la información necesaria para predecir el año siguiente.")
    anchor = insert_p_after(anchor, "Entre los algoritmos de aprendizaje automático, Random Forest presentó el mejor desempeño general, con RMSE de 3.789,03 mil litros, MAE de 1.448,93 mil litros, R² de 0,9675 y MAPE de 22,64%. Aunque no superó el baseline de persistencia, el valor elevado de R² demuestra que el modelo capturó de forma consistente la estructura productiva de los municipios. El XGBoost obtuvo RMSE de 6.341,84 mil litros, R² de 0,9088 y MAPE de 24,12%, indicando mayor sensibilidad a la configuración de hiperparámetros y a la escala desigual entre municipios.")
    anchor = insert_p_after(anchor, "La comparación con el modelo de persistencia es importante porque evita una interpretación exagerada de los algoritmos complejos. En problemas agrícolas con fuerte continuidad temporal, como la producción anual de leche, un baseline simple puede ser competitivo. Así, el resultado no invalida el uso de aprendizaje automático; al contrario, muestra que modelos como Random Forest y XGBoost deben ser comparados con referencias simples y ampliados con variables explicativas más ricas para justificar su aplicación operacional.")

    for fig, cap, desc in [
        ("figura_01_rmse_modelos.png", "Figura 1. Comparación del RMSE entre los modelos evaluados.", "La Figura 1 evidencia la diferencia de error cuadrático entre los tres enfoques. El menor RMSE fue observado en la persistencia lag-1, seguida por Random Forest. La distancia entre Random Forest y XGBoost sugiere que, con la configuración utilizada, el modelo de boosting fue menos estable para generalizar el comportamiento de los municipios en el período de prueba."),
        ("figura_02_observado_predicho.png", "Figura 2. Relación entre producción observada y producción predicha por los modelos.", "La Figura 2 permite evaluar visualmente la proximidad entre valores observados y predichos. Puntos más próximos a la diagonal indican predicciones más acertadas. La concentración de puntos cerca de la diagonal confirma el alto R² obtenido, pero las dispersiones en municipios de producción elevada muestran dónde los errores absolutos se vuelven más relevantes."),
        ("figura_03_series_municipios_lideres.png", "Figura 3. Series históricas de municipios líderes en producción lechera.", "La Figura 3 contextualiza la dinámica temporal de los municipios líderes. La estabilidad de las trayectorias explica por qué las variables rezagadas tuvieron alto poder predictivo. También muestra que municipios con gran escala productiva pueden dominar el error absoluto, reforzando la necesidad de analizar simultáneamente RMSE, MAE y MAPE."),
        ("figura_04_importancia_variables.png", "Figura 4. Importancia de variables del mejor modelo de aprendizaje automático.", "La Figura 4 muestra que la producción del año anterior fue la variable más influyente. Esto confirma que el componente histórico de la serie fue determinante. Las variables climáticas presentaron menor importancia relativa, probablemente por la agregación anual y por el uso de una única estación como señal regional."),
    ]:
        anchor = insert_picture_after(anchor, FIGS / fig, cap)
        anchor = insert_p_after(anchor, desc)

    imp = pd.read_csv(IMPORTANCE).head(10)
    imp_rows = [["Variable", "Importancia"]]
    for _, row in imp.iterrows():
        imp_rows.append([row["variable"], fmt(row["importancia"], 4)])
    _, anchor = insert_table_after(anchor, imp_rows, font_size=8)
    anchor = insert_p_after(anchor, "Tabla 6. Diez variables más importantes en el modelo Random Forest.", size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)

    for text in [
        "La variable más importante fue leite_lag_1, con importancia de 0,9564, seguida por leite_media_3anos y leite_lag_2. Esto confirma que la dinámica histórica de producción explica la mayor parte del comportamiento futuro. En términos prácticos, municipios que ya poseen trayectoria estable de producción tienden a mantener su posición relativa en el corto plazo, especialmente cuando la predicción se realiza en escala anual.",
        "Las variables climáticas aparecieron con importancia menor, lo que puede estar relacionado con tres factores: la agregación anual reduce eventos extremos de corta duración; la estación A819 de Castro no representa toda la variabilidad espacial del Paraná; y la producción municipal de leche depende también de manejo, alimentación, rebaño, tecnología, mercado y políticas públicas, variables que no fueron incorporadas en esta versión.",
        "Desde el punto de vista agrícola, el resultado es relevante porque muestra que bases públicas pueden apoyar estimaciones municipales de producción lechera con bajo costo computacional. Un sistema de este tipo puede auxiliar cooperativas, gestores públicos e industrias lácteas en diagnósticos territoriales, planeamiento de demanda, identificación de municipios con mayor estabilidad productiva y construcción de escenarios preliminares para políticas de apoyo a la cadena de leche.",
        "La comparación también revela un punto metodológico importante: cuando la serie histórica es muy fuerte, un algoritmo complejo no necesariamente supera un baseline temporal simple. Por eso, para estudios futuros, se recomienda ampliar la base con información espacialmente distribuida, como precipitación CHIRPS, cobertura de pasturas MapBiomas, relevo TOPODATA y características socioeconómicas o zootécnicas. Con esta ampliación, modelos más avanzados, incluyendo aprendizaje profundo y técnicas de explicabilidad, pueden tener mayor potencial de contribuir con predicciones robustas.",
        "Como aplicación futura, el modelo puede evolucionar hacia un sistema de apoyo a la decisión para agricultura de precisión territorial. La integración con bases geoespaciales permitiría identificar regiones con mayor riesgo climático, estimar impactos de sequías o exceso de lluvia y comparar escenarios de producción municipal. Para una tesis de maestría, el principal avance sería pasar de una predicción tabular simplificada hacia una fusión de datos agropecuarios, climáticos y geoespaciales con explicabilidad.",
    ]:
        anchor = insert_p_after(anchor, text)

    # Light formatting pass for top-level existing headings and title.
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if not txt:
            continue
        if i == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_font(r, size=11, bold=True)
        elif txt in ["RESUMEN", "PALABRAS CLAVE", "ABSTRACT", "KEYWORDS", "INTRODUCCIÓN", "MATERIALES Y MÉTODOS", "RESULTADOS Y DISCUSIÓN", "CONCLUSIÓN", "REFERENCIAS BIBLIOGRÁFICAS"]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_font(r, size=10, bold=True)
        elif txt.startswith("1.") or txt in ["Diseño metodológico", "Archivos utilizados y reproducibilidad", "Bases de datos realmente utilizadas", "Construcción de la base integrada", "Creación de los algoritmos"]:
            for r in p.runs:
                set_font(r, size=10, bold=True)
            p.paragraph_format.first_line_indent = Cm(0)
        else:
            for r in p.runs:
                if r.text:
                    set_font(r, size=10)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
