from __future__ import annotations

import hashlib
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\gabri\Downloads\Artigo_IIIWCAA_NOME.docx")
OUT = ROOT / "docs" / "Artigo_IIIWCAA_Alerta_Quedas_Leite.docx"
METRICS_PATH = ROOT / "results" / "tables" / "metricas_modelos.csv"
PRED_PATH = ROOT / "results" / "tables" / "predicciones_prueba.csv"
IMPORTANCE_PATH = ROOT / "results" / "tables" / "importancia_variables.csv"
BASE_PATH = ROOT / "data" / "base_integrada_leite_agroclima.csv"
FIG = ROOT / "results" / "figures"
ALERT_DIR = ROOT / "results" / "decline_risk"


def fmt(x: float, digits: int = 2) -> str:
    return f"{x:,.{digits}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def run_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def set_font(run, size=9, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_text(doc, text="", *, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=9.5, after=4, before=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, italic=italic)
    return p


def add_heading(doc, text):
    p = add_text(doc, text, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, size=10, before=3, after=2)
    p.paragraph_format.keep_with_next = True
    return p


def add_table(doc, headers, rows, widths=None, font_size=7.2, caption=None):
    table = doc.add_table(rows=2 if caption else 1, cols=len(headers))
    table.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "AAB8B8")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    cell_margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", 20), ("bottom", 20), ("left", 40), ("right", 40)):
        margin = OxmlElement(f"w:{edge}")
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")
        cell_margins.append(margin)
    table._tbl.tblPr.append(cell_margins)
    header_index = 0
    if caption:
        caption_cell = table.rows[0].cells[0].merge(table.rows[0].cells[-1])
        caption_cell.text = caption
        p = caption_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        for r in p.runs:
            set_font(r, 6.8, italic=True)
        header_index = 1
    for i, value in enumerate(headers):
        cell = table.rows[header_index].cells[i]
        cell.text = str(value)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                set_font(r, font_size, bold=True)
                r.font.color.rgb = RGBColor(31, 78, 77)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_font(r, font_size)
    for index, row in enumerate(table.rows):
        row_props = row._tr.get_or_add_trPr()
        row_props.append(OxmlElement("w:cantSplit"))
        if index == header_index:
            repeat_header = OxmlElement("w:tblHeader")
            repeat_header.set(qn("w:val"), "true")
            row_props.append(repeat_header)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc, path: Path, caption: str):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(3.05))
    c = add_text(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, size=8, after=4)
    c.paragraph_format.keep_together = True


def make_single_page_figure(metrics: pd.DataFrame, path: Path):
    ordered = metrics.sort_values("RMSE", ascending=True)
    fig, ax = plt.subplots(figsize=(5.2, 2.2))
    colors = ["#2a9d8f" if "Persistencia" in name else "#457b9d" if "Random Forest" in name else "#e76f51" for name in ordered["modelo"]]
    ax.barh(ordered["modelo"], ordered["RMSE"], color=colors, height=0.62)
    ax.set_xlabel("RMSE (mil litros)", fontsize=8)
    ax.set_title("Erro no teste temporal (2017-2019)", fontsize=9)
    ax.tick_params(axis="both", labelsize=7)
    ax.grid(axis="x", alpha=0.25)
    ax.set_axisbelow(True)
    fig.tight_layout(pad=0.6)
    fig.savefig(path, dpi=200, bbox_inches="tight")
    plt.close(fig)


def build_one_page_body(doc, metrics, base, predictions):
    rf = metrics.loc[metrics["modelo"] == "Random Forest"].iloc[0]
    xgb = metrics.loc[metrics["modelo"] == "XGBoost"].iloc[0]
    baseline = metrics.loc[metrics["modelo"] == "Sem queda (baseline)"].iloc[0]
    prevalence = float(predictions.drop_duplicates(["codigo_municipio", "ano"])["queda_10_pct"].mean())

    add_text(doc, f"Resumo: Este estudo investiga um problema distinto da previsão de volumes: identificar municípios com queda anual de produção de leite igual ou superior a 10%. A partir da PPM/IBGE-SIDRA (tabela 74, produto 2682) e do painel agroclimático público já integrado, foram usados somente atributos disponíveis até t-1. Floresta Aleatória e XGBoost foram avaliados em teste temporal 2017-2019 (1.197 observações); a prevalência de quedas foi {fmt(prevalence*100,1)}%. Os resultados ficaram próximos da referência sem queda, indicando que esses dados não sustentaram um alerta confiável neste recorte.", size=8.5, after=3)
    add_text(doc, "Palavras-chave: produção leiteira; classificação; alerta de queda; floresta aleatória; XGBoost.", size=8.3, after=3)

    add_heading(doc, "1. Introdução")
    add_text(doc, "A aprendizagem de máquina tem sido aplicada à predição da produção e ao manejo leiteiro, em tarefas cuja escala e desempenho dependem dos dados disponíveis (Slob et al., 2021; Palma et al., 2025). Estudos modelaram séries mensais de leite em sistemas de pastagem (Perez-Guerra et al., 2023) e produção nacional com dados climáticos e econômicos (Cesarini et al., 2024). Na agricultura brasileira, clima e sensoriamento remoto também foram integrados para prever produtividade municipal de soja (von Bloh et al., 2023). Esses trabalhos estimam volumes; este artigo propõe um recorte diferente: classificar retrações anuais da produção leiteira municipal, sem atribuir causas, e avaliar se dados públicos defasados permitem sinalizá-las.", size=8.4, after=3)

    add_heading(doc, "2. Materiais e Métodos")
    add_text(doc, f"A unidade município-ano abrange {base['codigo_municipio'].nunique()} municípios. A produção anual (mil litros) vem da PPM/IBGE-SIDRA, tabela 74, produto leite 2682 (https://sidra.ibge.gov.br/tabela/74). O alvo vale 1 quando a queda ante t-1 é de pelo menos 10%, limiar operacional. Preditores: produção em t-1/t-2, crescimento até t-1, ano, código municipal categórico e clima defasado da estação INMET A819, em Castro (https://portal.inmet.gov.br/dadoshistoricos): precipitação, temperaturas média e máxima média, radiação, umidade, vento e dias secos aproximados. O clima de t não prevê o próprio t; após remover ausências necessárias aos lags, modelaram-se 2008-2019. O clima de Castro é comum aos municípios, e o alvo classifica quedas, não volume.", size=8.2, after=3)
    add_text(doc, "Compararam-se Floresta Aleatória (RF) e XGBoost (XGB) com a referência ‘sem queda’. A precisão média (AP) selecionou hiperparâmetros em validações cronológicas de 2015 e 2016, treinadas apenas com anos anteriores; 2017-2019 ficou como teste intocado. No RF, a busca combinou 200/500 árvores, profundidade 6/12 e folha mínima 2/5, com pesos balanceados; selecionou 500, 12 e 5. No XGB, testaram-se 150/300 árvores, profundidade 2/4, taxa 0,03/0,08 e subamostragens de 0,8; selecionou 150, 2 e 0,08. Usou-se limiar 0,5. Métricas: precisão, sensibilidade, F1, acurácia balanceada, AP e área sob a curva ROC (ROC-AUC). Código: https://github.com/GabyQueiroz/CITREN_leite, script scripts/build_decline_risk_models.py.", size=8.2, after=2)

    add_heading(doc, "3. Resultados")
    metric_rows = []
    for model in ("Sem queda (baseline)", "Random Forest", "XGBoost"):
        row = metrics.loc[metrics["modelo"] == model].iloc[0]
        label = "Sem queda (referência)" if model == "Sem queda (baseline)" else "Floresta Aleatória" if model == "Random Forest" else model
        metric_rows.append([label, fmt(row["precision"],2), fmt(row["recall"],2), fmt(row["f1"],2), fmt(row["average_precision"],2), fmt(row["roc_auc"],2)])
    add_table(doc, ["Modelo", "Prec.", "Sens.", "F1", "AP", "ROC-AUC"], metric_rows, widths=[1.0, 0.4, 0.45, 0.35, 0.35, 0.4], font_size=6.2, caption="Tabela 1. Métricas de classificação no teste temporal (2017-2019). Fonte: elaboração própria.")
    add_figure(doc, ALERT_DIR / "figures" / "matrices_confusion_test.png", "Figura 1. Casos reais e sinalizados pelos classificadores com limiar 0,5. Fonte: elaboração própria.")
    add_text(doc, f"As quedas corresponderam a {fmt(prevalence*100,1)}% do teste. A Floresta Aleatória identificou 85 dos 289 casos (sensibilidade={fmt(rf['recall'],2)}), mas gerou 296 falsos alertas e precisão de {fmt(rf['precision'],2)}; sua AP ({fmt(rf['average_precision'],3)}) ficou abaixo da referência ({fmt(baseline['average_precision'],3)}). O XGBoost teve AP={fmt(xgb['average_precision'],3)}, próxima à prevalência: não identificou quedas e gerou três falsos alertas. As acurácias balanceadas foram {fmt(rf['balanced_accuracy'],3)} (RF) e {fmt(xgb['balanced_accuracy'],3)} (XGB); ROC-AUCs próximas de 0,5 indicam discriminação insuficiente. Os resultados não justificam uso operacional como alerta.", size=8.1, after=3)
    add_text(doc, "A regra ‘sem queda’ acertaria a maioria das observações por causa do desbalanceamento, mas teria sensibilidade nula; por isso, acurácia simples não seria evidência de alerta útil. O RF recuperou menos de um terço das retrações e mais de três quartos dos seus avisos foram falsos. A AP próxima da prevalência mostra que ordenar municípios por risco quase não melhorou em relação ao acaso. Em aplicações reais, falsos alertas também têm custo operacional e exigiriam validação com agentes do setor.", size=8.1, after=3)

    add_heading(doc, "4. Conclusão")
    add_text(doc, "Este recorte formula uma pergunta diferente da regressão do volume anual: avaliar detecção de retrações com atributos defasados e teste cronológico. Neste experimento, RF e XGBoost não superaram a referência simples; portanto, não se recomenda implantar alertas com estas variáveis. A PPM anual e o clima de uma única estação, replicado no estado, oferecem pouca informação espacial, e a queda de 10% é uma definição operacional sensível a revisões e flutuações. Estudos futuros devem incluir clima espacializado, tamanho e estrutura do rebanho e séries mais longas, validar o limiar com especialistas e avaliar sensibilidade, calibração e validação externa.", size=8.1, after=3)

    add_heading(doc, "Referências")
    refs = [
        "IBGE. Pesquisa da Pecuária Municipal, tabela SIDRA 74. https://sidra.ibge.gov.br/tabela/74. Acesso em: 20 out. 2026.",
        "INMET. Dados históricos de estações meteorológicas. https://portal.inmet.gov.br/dadoshistoricos. Acesso em: 20 out. 2026.",
        "CESARINI, L. et al. Comparison of deep learning models for milk production forecasting at national scale. Computers and Electronics in Agriculture, 221, 108933, 2024. doi:10.1016/j.compag.2024.108933.",
        "PALMA, O. et al. AI and data analytics in the dairy farms: a scoping review. Animals, 15(9), 1291, 2025. doi:10.3390/ani15091291.",
        "PEREZ-GUERRA, U. H. et al. Seasonal autoregressive integrated moving average (SARIMA) time-series model for milk production forecasting in pasture-based dairy cows in the Andean highlands. PLOS ONE, 18, e0288849, 2023. doi:10.1371/journal.pone.0288849.",
        "SLOB, N.; CATAL, C.; KASSAHUN, A. Application of machine learning to improve dairy farm management. Preventive Veterinary Medicine, 187, 105237, 2021. doi:10.1016/j.prevetmed.2020.105237.",
        "VON BLOH, M. et al. Machine learning for soybean yield forecasting in Brazil. Agricultural and Forest Meteorology, 341, 109670, 2023. doi:10.1016/j.agrformet.2023.109670.",
    ]
    for ref in refs:
        add_text(doc, ref, align=WD_ALIGN_PARAGRAPH.LEFT, size=6.5, after=1)


def main():
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Modelo não encontrado: {TEMPLATE}")
    metrics_path = ALERT_DIR / "tables" / "metricas_alerta_queda.csv"
    predictions_path = ALERT_DIR / "tables" / "predicoes_alerta_queda.csv"
    if not metrics_path.exists() or not predictions_path.exists():
        raise FileNotFoundError("Execute scripts/build_decline_risk_models.py antes de gerar o artigo.")

    metrics = pd.read_csv(metrics_path)
    pred = pd.read_csv(predictions_path)
    base = pd.read_csv(BASE_PATH)
    source_hash = run_hash(TEMPLATE)

    doc = Document(TEMPLATE)
    # Normalize fractional twips in this template so python-docx can calculate table widths.
    for section in doc.sections:
        margins = section._sectPr.find(qn("w:pgMar"))
        if margins is not None:
            for key, value in list(margins.attrib.items()):
                try:
                    margins.set(key, str(round(float(value))))
                except ValueError:
                    pass
    body = doc._element.body
    existing = list(body)
    # Retain the title/author area and its section break; remove only model prose below it.
    for child in existing[4:-1]:
        body.remove(child)

    doc.paragraphs[0].text = "CLASSIFICAÇÃO TEMPORAL DE QUEDAS NA PRODUÇÃO MUNICIPAL DE LEITE NO PARANÁ COM APRENDIZADO DE MÁQUINA"
    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in doc.paragraphs[0].runs:
        set_font(run, 12, bold=True)
    authors = [
        "Andressa Aparecida Guimarães - Universidade Estadual de Ponta Grossa (UEPG)",
        "Gabrielly de Queiroz Pereira - Universidade Estadual de Ponta Grossa (UEPG)",
        "Maria Salete Marcon Gomes Vaz - Universidade Estadual de Ponta Grossa (UEPG)",
    ]
    for i, text in enumerate(authors, start=1):
        doc.paragraphs[i].text = text
        doc.paragraphs[i].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in doc.paragraphs[i].runs:
            set_font(run, 8.5, bold=True)

    build_one_page_body(doc, metrics, base, pred)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Artigo de uma página criado: {OUT}")


if __name__ == "__main__":
    main()
